from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.models.contact import Contact
from app.schemas.contact import ContactCreate
from app.core.auth import require_admin
from app.core.email import (
    send_contact_notification,
    send_client_confirmation,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO


router = APIRouter(prefix="/api/contact", tags=["Contact"])


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# =========================================================
# CREATE CLIENT INQUIRY
# =========================================================

@router.post("/")
async def create_contact(
    contact: ContactCreate,
    db: Session = Depends(get_db),
):
    new_contact = Contact(
        name=contact.name,
        email=contact.email,
        phone=contact.phone,
        city=contact.city,
        property_type=contact.property_type,
        project_type=contact.project_type,
        budget=contact.budget,
        message=contact.message,
    )

    db.add(new_contact)
    db.commit()
    db.refresh(new_contact)

    # =====================================================
    # SEND EMAIL NOTIFICATIONS
    # =====================================================

    # -----------------------------------------------------
    # Email to Tara Living
    # -----------------------------------------------------

    try:
        await send_contact_notification(
            name=contact.name,
            email=contact.email,
            phone=contact.phone,
            city=contact.city,
            property_type=contact.property_type,
            project_type=contact.project_type,
            budget=contact.budget,
            message=contact.message,
        )

        print("✅ Owner notification sent")

    except Exception as e:
        print("❌ Owner email failed:", e)

    # -----------------------------------------------------
    # Confirmation email to client
    # -----------------------------------------------------

    try:
        await send_client_confirmation(
            name=contact.name,
            email=contact.email,
            property_type=contact.property_type,
            project_type=contact.project_type,
            city=contact.city,
        )

        print("✅ Client confirmation sent")

    except Exception as e:
        print("❌ Client confirmation failed:", e)

    return {
        "message": "Inquiry submitted successfully",
        "id": new_contact.id,
    }


# =========================================================
# GET ALL CLIENT INQUIRIES
# =========================================================

@router.get("/")
def get_contacts(
    db: Session = Depends(get_db),
    admin=Depends(require_admin),
):
    contacts = (
        db.query(Contact)
        .order_by(Contact.created_at.desc())
        .all()
    )

    return contacts


# =========================================================
# UPDATE CLIENT STATUS
# =========================================================

@router.patch("/{contact_id}/status")
def update_contact_status(
    contact_id: int,
    status: str,
    db: Session = Depends(get_db),
    admin=Depends(require_admin),
):
    allowed_statuses = [
        "New",
        "Contacted",
        "Site Visit",
        "Completed",
        "Deal Not Done",
    ]

    if status not in allowed_statuses:
        raise HTTPException(
            status_code=400,
            detail="Invalid client status",
        )

    contact = (
        db.query(Contact)
        .filter(Contact.id == contact_id)
        .first()
    )

    if not contact:
        raise HTTPException(
            status_code=404,
            detail="Client not found",
        )

    contact.status = status

    db.commit()
    db.refresh(contact)

    return {
        "message": "Client status updated successfully",
        "id": contact.id,
        "status": contact.status,
    }

# =========================================================
# DELETE CLIENT INQUIRY
# =========================================================

@router.delete("/{contact_id}")
def delete_contact(
    contact_id: int,
    db: Session = Depends(get_db),
    admin=Depends(require_admin),
):
    contact = (
        db.query(Contact)
        .filter(Contact.id == contact_id)
        .first()
    )

    if not contact:
        raise HTTPException(
            status_code=404,
            detail="Client not found",
        )

    db.delete(contact)
    db.commit()

    return {
        "message": "Client deleted successfully",
        "id": contact_id,
    }


# =========================================================
# EXPORT CLIENT INQUIRIES TO EXCEL
# =========================================================

@router.get("/export")
def export_contacts_to_excel(
    db: Session = Depends(get_db),
    admin=Depends(require_admin),
):
    contacts = (
        db.query(Contact)
        .order_by(Contact.created_at.desc())
        .all()
    )

    # Create Excel workbook
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Client Inquiries"

    # Excel headers
    headers = [
        "ID",
        "Name",
        "Email",
        "Phone",
        "City",
        "Property Type",
        "Project Type",
        "Budget",
        "Message",
        "Status",
        "Created At",
    ]

    worksheet.append(headers)

    # Header styling
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="24231F",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Add client data
    for index, contact in enumerate(contacts, start=1):
        worksheet.append([
            index,
            contact.name,
            contact.email,
            contact.phone,
            contact.city,
            contact.property_type,
            contact.project_type,
            contact.budget,
            contact.message,
            contact.status,
            (
                contact.created_at.strftime("%d-%b-%Y %H:%M")
                if contact.created_at
                else ""
            ),
        ])

    # Column widths
    column_widths = {
        "A": 8,
        "B": 25,
        "C": 32,
        "D": 18,
        "E": 20,
        "F": 22,
        "G": 25,
        "H": 20,
        "I": 50,
        "J": 15,
        "K": 22,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # Make message cells wrap
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        row[8].alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Create downloadable Excel file in memory
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; '
                'filename="tara-living-client-inquiries.xlsx"'
            )
        },
    )